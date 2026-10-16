"""
generate_daily_breadth_report.py (USA — NASDAQ + NYSE, equities only)

Builds a day-by-day US market breadth summary for a lookback period and
saves it to an Excel file. Reuses compute_live_breadth() from utils/live.py
UNCHANGED — the same function the app's Live tab uses for "right now" —
applied here once per historical trading day instead.

compute_live_breadth() reads a "% Change" column to derive up_4pct,
down_4pct, advancers, and decliners. For each historical day, this script
reconstructs that same shape from Close-to-Close % change across the whole
loaded universe, then hands it to compute_live_breadth() exactly as if it
were a live snapshot for that day.

Output columns: Date, Up 4%+, Down 4%+, Advancers, Decliners

*** IMPORTANT — runtime warning ***
The full NASDAQ+NYSE equities universe is roughly 6,000-8,000 symbols.
Downloading 3 months of daily history for ALL of them via free yfinance
(sequential downloads, by design — see the main app's README on why
threads=False) can take a very long time, potentially HOURS, and free
Yahoo Finance may rate-limit a run that large. Two ways to make this
practical:
  1. Set MIN_PRICE / MAX_PRICE below to narrow the universe first (a cheap
     last-price pass runs before the slow full-history download, same
     two-stage approach the main app uses for its full-market scan).
  2. Just let it run unattended (e.g. overnight) if you want the true
     full, unfiltered market and don't mind the wait.

Run:
    python generate_daily_breadth_report.py

Requires network access to Yahoo Finance via yfinance — run this locally,
not in a restricted sandbox.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd
from openpyxl.styles import Font, Alignment

from utils.data import (
    fetch_full_listed_universe, filter_equities_only,
    fetch_last_close_prices, fetch_price_history,
)
try:
    from utils.data import prices_from_listing
except ImportError:  # older utils/data.py (no list-supplied prices): every price is fetched from Yahoo, as before
    def prices_from_listing(universe):
        return pd.Series(dtype=float)
from utils.live import compute_live_breadth

# ---------------- Configuration ----------------
LOOKBACK_PERIOD = "3mo"     # yfinance period string: "1mo", "3mo", "6mo", "1y", etc.
MIN_PRICE = None            # e.g. 1.0 — set to narrow the universe and speed this up a lot
MAX_PRICE = None            # e.g. 60.0 — matches the main app's default full-scan price cap
OUTPUT_PATH = "usa_daily_breadth_summary.xlsx"
CHUNK_SIZE = 150
PAUSE_SEC = 1.0


def build_daily_breadth_table(price_data: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """
    Reconstructs a day-by-day breadth summary by calling compute_live_breadth()
    once per historical trading day, exactly as the Live tab calls it for
    the current moment.
    """
    # Wide matrix: rows = date, columns = ticker, values = Close. Building a
    # DataFrame from a dict of Series aligns on the union of all dates
    # automatically, so tickers missing a bar on a given day just show NaN there.
    close_matrix = pd.DataFrame({ticker: df["Close"] for ticker, df in price_data.items()})
    close_matrix = close_matrix.sort_index()

    # Day-over-day % change per ticker, vectorized.
    pct_change_matrix = close_matrix.pct_change() * 100

    rows = []
    for dt, row in pct_change_matrix.iterrows():
        changes = row.dropna()
        if changes.empty:
            continue  # first day in the window (no prior close) or a non-trading day

        snapshot = pd.DataFrame({"% Change": changes})
        stats = compute_live_breadth(snapshot)

        rows.append({
            "Date": dt.date() if hasattr(dt, "date") else dt,
            "Up 4%+": stats["up_4pct"],
            "Down 4%+": stats["down_4pct"],
            "Advancers": stats["advancers"],
            "Decliners": stats["decliners"],
        })

    return pd.DataFrame(rows).sort_values("Date").reset_index(drop=True)


def save_to_excel(df: pd.DataFrame, path: str) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Daily Breadth")
        ws = writer.sheets["Daily Breadth"]

        header_font = Font(name="Arial", bold=True)
        body_font = Font(name="Arial")

        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font

        widths = {"A": 14, "B": 10, "C": 11, "D": 12, "E": 12}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

        for cell in list(ws["A"])[1:]:
            cell.number_format = "yyyy-mm-dd"
            cell.alignment = Alignment(horizontal="center")
        for col in ["B", "C", "D", "E"]:
            for cell in list(ws[col])[1:]:
                cell.alignment = Alignment(horizontal="center")


def main():
    print("Fetching NASDAQ + NYSE symbol directory...")
    raw_universe = fetch_full_listed_universe()
    if raw_universe.empty:
        print("Could not load a NASDAQ/NYSE symbol list from any source (see the error above). Try again shortly.")
        return

    equities = filter_equities_only(raw_universe)
    print(f"{len(raw_universe)} total listed symbols, {len(equities)} after equity filtering "
          f"(excluding ETFs/warrants/units/preferreds/funds — heuristic).")

    tickers = equities["YF_Symbol"].tolist()

    if MIN_PRICE is not None or MAX_PRICE is not None:
        lo_label = f"${MIN_PRICE:.2f}" if MIN_PRICE is not None else "$0"
        hi_label = f"${MAX_PRICE:.2f}" if MAX_PRICE is not None else "no cap"
        print(f"Applying price filter ({lo_label}-{hi_label}) before the full history download...")
        # Prices carried by the symbol list itself need no download; only symbols without one are fetched.
        listing_prices = prices_from_listing(equities)
        missing = [t for t in tickers if t not in listing_prices.index]
        print(f"{len(listing_prices)} symbols priced straight from the list; fetching {len(missing)} more from Yahoo...")
        fetched = fetch_last_close_prices(missing, chunk_size=200, pause_sec=PAUSE_SEC) if missing else {}
        price_series = pd.concat([listing_prices, pd.Series(fetched, dtype=float)])
        lo = MIN_PRICE if MIN_PRICE is not None else 0
        hi = MAX_PRICE if MAX_PRICE is not None else float("inf")
        tickers = price_series[(price_series >= lo) & (price_series <= hi)].index.tolist()
        print(f"{len(tickers)} symbols remain after price filtering.")
    else:
        print(f"WARNING: no price filter set — about to download {LOOKBACK_PERIOD} history for "
              f"{len(tickers)} symbols (the full NASDAQ+NYSE equities universe). This WILL take a "
              f"long time (potentially hours) on free yfinance. Set MIN_PRICE/MAX_PRICE at the top "
              f"of this script to narrow it down if that's impractical for you.")

    if not tickers:
        print("No tickers left to fetch after filtering.")
        return

    print(f"Downloading {LOOKBACK_PERIOD} price history for {len(tickers)} symbols "
          f"(chunked at {CHUNK_SIZE}/chunk, sequential — this is the slow part)...")
    price_data = fetch_price_history(tuple(tickers), period=LOOKBACK_PERIOD,
                                      chunk_size=CHUNK_SIZE, pause_sec=PAUSE_SEC)
    print(f"Loaded history for {len(price_data)} symbols.")

    if not price_data:
        print("No price data returned.")
        return

    print("Computing day-by-day breadth summary via compute_live_breadth()...")
    daily_df = build_daily_breadth_table(price_data)

    if daily_df.empty:
        print("No daily breadth data could be computed.")
        return

    save_to_excel(daily_df, OUTPUT_PATH)
    print(f"Saved {len(daily_df)} trading days to {OUTPUT_PATH}")
    print(daily_df.tail(5).to_string(index=False))


if __name__ == "__main__":
    main()
